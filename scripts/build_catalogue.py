import json
from datetime import UTC, date, datetime
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE_DIR = ROOT.parent / "catalogues" / "WWE collector app" / "Completed set catalogues"
OUTPUT = ROOT / "data" / "catalogue.json"

SET_CONFIG = {
    "2025_Topps_Universe_WWE_Master_Catalogue.xlsx": {
        "id": "2025-topps-universe-wwe",
        "name": "2025 Topps Universe WWE",
        "shortName": "Universe WWE",
        "accent": "#8b5cf6",
    },
    "2025_Topps_WWE_x_BAPE_Master_Catalogue.xlsx": {
        "id": "2025-topps-wwe-x-bape",
        "name": "2025 Topps WWE x BAPE",
        "shortName": "WWE x BAPE",
        "accent": "#f59e0b",
    },
    "2025_Topps_Finest_WWE_Master_Catalogue.xlsx": {
        "id": "2025-topps-finest-wwe",
        "name": "2025 Topps Finest WWE",
        "shortName": "Finest WWE",
        "accent": "#06b6d4",
    },
    "2025_Topps_Chrome_WWE_x_Cactus_Jack_Master_Catalogue.xlsx": {
        "id": "2025-topps-chrome-cactus-jack-wrestlemania",
        "name": "2025 Topps Chrome WWE Cactus Jack x WrestleMania",
        "shortName": "Cactus Jack x WrestleMania",
        "accent": "#ef4444",
        "releaseDate": "2025-04-16",
        "wrestlemaniaEdition": True,
    },
}

CACTUS_CHAMPIONS = [
    ("CCA-1", "John Cena", "WWE"),
    ("CCA-2", "Rhea Ripley", "Raw"),
    ("CCA-3", "Roman Reigns", "Smackdown"),
    ("CCA-4", "Jey Uso", "Raw"),
    ("CCA-5", "Gunther", "Raw"),
    ("CCA-6", "Damian Priest", "Raw"),
    ("CCA-7", "Randy Orton", "Smackdown"),
    ("CCA-8", "Bayley", "Raw"),
    ("CCA-9", "Oba Femi", "NXT"),
    ("CCA-10", 'Seth "Freakin" Rollins', "Raw"),
    ("CCA-11", "Roxanne Perez", "NXT"),
    ("CCA-12", "Drew McIntyre", None),
    ("CCA-13", "Nia Jax", "Smackdown"),
    ("CCA-14", "Trick Williams", "NXT"),
    ("CCA-15", "Kevin Owens", "Smackdown"),
    ("CCA-16", "Finn Bálor", "Raw"),
    ("CCA-17", "Liv Morgan", "Raw"),
    ("CCA-18", "CM Punk", "Raw"),
    ("CCA-19", "Becky Lynch", "WWE"),
    ("CCA-20", "LA Knight", "Smackdown"),
    ("CCA-21", "Tiffany Stratton", "Smackdown"),
    ("CCA-22", "Solo Sikoa", "Smackdown"),
    ("CCA-23", "Bianca Belair", "Smackdown"),
    ("CCA-24", "Sami Zayn", "Raw"),
    ("CCA-25", '"The American Nightmare" Cody Rhodes', "Smackdown"),
]

WRESTLEMANIA_AUTOGRAPHS = [
    ("CJA-BB", "Bianca Belair", "Smackdown"),
    ("CJA-BK", "Bron Breakker", "Raw"),
    ("CJA-CG", "Chelsea Green", "Smackdown"),
    ("CJA-CM", "CM Punk", "Raw"),
    ("CJA-CR", '"The American Nightmare" Cody Rhodes', "Smackdown"),
    ("CJA-DD", '"Dirty" Dominik Mysterio', "Raw"),
    ("CJA-DM", "Drew McIntyre", "Raw"),
    ("CJA-DP", "Damian Priest", "Raw"),
    ("CJA-GL", "Giulia", "NXT"),
    ("CJA-GT", "Gunther", "Raw"),
    ("CJA-JC", "John Cena", "WWE"),
    ("CJA-JU", "Jey Uso", "Raw"),
    ("CJA-LA", "LA Knight", "Smackdown"),
    ("CJA-LV", "Lyra Valkyria", "Raw"),
    ("CJA-PT", "Penta", "Raw"),
    ("CJA-RM", "Roman Reigns", "Smackdown"),
    ("CJA-RR", "Rhea Ripley", "Raw"),
    ("CJA-SR", 'Seth "Freakin" Rollins', "Raw"),
    ("CJA-SS", "Solo Sikoa", "Smackdown"),
    ("CJA-SV", "Stephanie Vaquer", "NXT"),
    ("CJA-TR", "The Rock", "Legend"),
    ("CJA-TS", "Tiffany Stratton", "Smackdown"),
    ("CJA-TV", "Travis Scott", None),
    ("CJA-UT", "Undertaker", "Legend"),
    ("CJA-ZR", "Zaria", "NXT"),
]


def clean(value):
    if value is None:
        return None
    if isinstance(value, (date, datetime)):
        return value.isoformat()[:10]
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def rows_as_dicts(sheet):
    rows = sheet.iter_rows(values_only=True)
    headers = [clean(value) for value in next(rows)]
    for row in rows:
        values = [clean(value) for value in row]
        if not any(value is not None for value in values):
            continue
        yield dict(zip(headers, values))


def card_from_row(row, set_id):
    return {
        "id": row["Card UID"],
        "setId": set_id,
        "order": int(row["Checklist Order"]),
        "number": str(row["Card Number"]),
        "name": row["Display Name"],
        "subject2": row["Subject 2"],
        "category": row["Category"],
        "subset": row["Subset"],
        "subsetCode": row["Subset Code"],
        "roster": row["Roster"],
        "rookie": row["Rookie"],
        "parallelGroup": row["Parallel Group"],
    }


def build_wrestlemania_cards(master_rows, set_id):
    base_rows = [row for row in master_rows if row["Category"] == "Base"][:100]
    base_cards = [card_from_row(row, set_id) for row in base_rows]

    inserts = [
        {
            "id": f"2025-WMCJ-{number}",
            "setId": set_id,
            "order": 100 + index,
            "number": number,
            "name": name,
            "subject2": None,
            "category": "Insert",
            "subset": "Cactus Champions",
            "subsetCode": "CCA",
            "roster": roster,
            "rookie": "No",
            "parallelGroup": "CACTUS_CHAMPIONS",
        }
        for index, (number, name, roster) in enumerate(CACTUS_CHAMPIONS, 1)
    ]

    autographs = [
        {
            "id": f"2025-WMCJ-{number}",
            "setId": set_id,
            "order": 125 + index,
            "number": number,
            "name": name,
            "subject2": None,
            "category": "Autograph",
            "subset": "Base Cards Autograph Variation",
            "subsetCode": "CJA",
            "roster": roster,
            "rookie": "No",
            "parallelGroup": "BASE_AUTOGRAPH",
        }
        for index, (number, name, roster) in enumerate(WRESTLEMANIA_AUTOGRAPHS, 1)
    ]
    return base_cards + inserts + autographs


def main():
    sets = []
    cards = []

    for filename, config in SET_CONFIG.items():
        path = SOURCE_DIR / filename
        workbook = load_workbook(path, read_only=True, data_only=True)

        master_rows = list(rows_as_dicts(workbook["Master Cards"]))
        if config.get("wrestlemaniaEdition"):
            set_cards = build_wrestlemania_cards(master_rows, config["id"])
            subsets = [
                {"category": "Base", "name": "Base Cards", "code": "BASE", "count": 100},
                {"category": "Insert", "name": "Cactus Champions", "code": "CCA", "count": 25},
                {"category": "Autograph", "name": "Base Cards Autograph Variation", "code": "CJA", "count": 25},
            ]
        else:
            set_cards = [card_from_row(row, config["id"]) for row in master_rows]
            subsets = [
                {
                    "category": row["Category"],
                    "name": row["Subset"],
                    "code": row["Subset Code"],
                    "count": int(row["Card Count"]),
                }
                for row in rows_as_dicts(workbook["Subsets"])
            ]

        cards.extend(set_cards)
        release_date = config.get("releaseDate") or (master_rows[0]["Release Date"] if master_rows else None)
        sets.append({
            **{key: value for key, value in config.items() if key not in {"releaseDate", "wrestlemaniaEdition"}},
            "year": 2025,
            "manufacturer": "Topps",
            "releaseDate": release_date,
            "cardCount": len(set_cards),
            "subsetCount": len(subsets),
            "subsets": subsets,
        })

    payload = {
        "schemaVersion": 1,
        "generatedAt": datetime.now(UTC).replace(microsecond=0).isoformat().replace("+00:00", "Z"),
        "setCount": len(sets),
        "cardCount": len(cards),
        "sets": sets,
        "cards": cards,
    }
    OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(f"Wrote {len(cards):,} cards across {len(sets)} sets to {OUTPUT}")


if __name__ == "__main__":
    main()
